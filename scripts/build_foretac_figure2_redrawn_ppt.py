#!/usr/bin/env python3
"""Build Figure 2 from a blank slide, with editable vector structure."""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

import build_foretac_figure2_ppt as base


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "paper/figures/figure2_chip_assets"
BOARD = ROOT / "paper/figures/board_marker_displacement/panels"
OUTPUT = ROOT / "paper/figures/ForeTac_Figure2_redrawn_v4_20260812.pptx"
DESKTOP = Path("/home/chenshuai/Desktop/ForeTac_Figure2_redrawn_v4_20260812.pptx")


def curved_arrow(slide, x1, y1, x2, y2, color, width=1.25, dashed=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.CURVE, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.name = "EDITABLE_CURVED_FLOW"
    base.line(shape, color, width)
    if dashed:
        base.set_dash(shape)
    base.arrow_end(shape)
    return shape


def elbow_arrow(slide, x1, y1, x2, y2, color, width=1.1, dashed=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.ELBOW, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.name = "EDITABLE_ELBOW_FLOW"
    base.line(shape, color, width)
    if dashed:
        base.set_dash(shape)
    base.arrow_end(shape)
    return shape


def sample_card(slide, path, x, y, positive, label):
    edge = base.C["green"] if positive else base.C["red"]
    base.picture_cover(slide, path, x, y, 0.38, 0.38, edge, 1.0)
    badge = "+" if positive else "-"
    base.box(
        slide, x + 0.29, y - 0.04, 0.14, 0.14, edge, base.C["white"],
        radius=1, width=0.7, label=badge, size=7.0, bold=True,
        text_color=base.C["white"],
    )
    base.textbox(slide, label, x - 0.04, y + 0.40, 0.46, 0.15, 5.4, base.C["muted"])


def draw_scorer(slide):
    base.textbox(slide, "(c) Contact-Quality Scorer", 10.20, 0.47, 2.95, 0.24, 10.2, base.C["ink"], True)
    examples = [
        (BOARD / "board_marker_contact_modes_stable_contact.png", True, "stable"),
        (ASSETS / "marker_rdp/chip_ep01_f0046_left_rdp.png", False, "dropout"),
        (BOARD / "board_marker_contact_modes_insufficient_pressure.png", False, "low force"),
        (BOARD / "board_marker_contact_modes_excessive_pressure.png", False, "excess"),
        (BOARD / "board_marker_contact_modes_oscillation.png", False, "slip"),
    ]
    for i, item in enumerate(examples):
        sample_card(slide, item[0], 10.28 + i * 0.56, 0.91, item[1], item[2])

    base.connector(slide, 10.26, 1.53, 13.05, 1.53, base.C["line"], 0.65, arrow=False)
    base.textbox(slide, "aligned action", 10.26, 1.65, 1.10, 0.16, 6.2, base.C["orange_dark"], True)
    base.token_row(slide, 10.34, 1.87, 5, base.C["orange"], 0.13, 0.035)
    base.textbox(slide, "predicted touch", 10.26, 2.12, 1.10, 0.16, 6.2, base.C["teal_dark"], True)
    base.token_row(slide, 10.34, 2.34, 5, base.C["cyan"], 0.13, 0.035, base.C["teal_dark"])

    base.box(
        slide, 11.55, 1.86, 0.78, 0.72, base.C["gold_light"], base.C["gold"],
        label="Contact-Quality\nScorer", size=7.0, bold=True,
    )
    curved_arrow(slide, 11.16, 1.94, 11.55, 2.08, base.C["orange"], 1.0)
    curved_arrow(slide, 11.16, 2.41, 11.55, 2.31, base.C["teal"], 1.0)

    base.textbox(slide, "stable contact", 12.45, 1.83, 0.63, 0.16, 5.8, base.C["green_dark"], True, PP_ALIGN.LEFT)
    base.box(slide, 12.46, 2.03, 0.53, 0.10, base.C["green"], base.C["green"], radius=0.01)
    base.textbox(slide, "failure risk", 12.45, 2.25, 0.63, 0.16, 5.8, base.C["red"], True, PP_ALIGN.LEFT)
    base.box(slide, 12.46, 2.45, 0.31, 0.10, base.C["red"], base.C["red"], radius=0.01)
    base.connector(slide, 12.33, 2.22, 12.44, 2.22, base.C["gold"], 1.0)


def draw_inference(slide, assets):
    base.add_section_title(slide, "INFERENCE-TIME PREDICTIVE GUIDANCE", 0.18, 3.24, 4.2)
    base.textbox(
        slide, "Frozen modules; guidance updates only the selected action state.",
        0.18, 3.47, 3.7, 0.16, 6.4, base.C["muted"], False, PP_ALIGN.LEFT,
    )

    # Compact predictive branch. Curved arrows distinguish conditioning from the main flow.
    base.token_row(slide, 3.15, 3.78, 3, base.C["blue"], 0.11, 0.025)
    base.latent_grid(slide, 3.55, 3.72, 0.065, base.C["teal_light"], base.C["teal"])
    base.textbox(slide, "current context", 3.07, 3.55, 0.80, 0.16, 6.0, base.C["blue_dark"], True)

    base.box(slide, 4.08, 3.68, 0.72, 0.50, base.C["blue_light"], base.C["blue_dark"], label="Denoiser", size=6.5, bold=True)
    base.add_lock(slide, 4.68, 3.58, 0.7)
    curved_arrow(slide, 3.86, 3.91, 4.08, 3.94, base.C["blue_dark"], 1.0)

    base.token_row(slide, 5.02, 3.87, 4, base.C["orange"], 0.12, 0.025)
    base.textbox(slide, "clean action estimate", 4.92, 3.57, 1.00, 0.18, 6.0, base.C["orange_dark"], True)
    curved_arrow(slide, 4.80, 3.94, 5.00, 3.94, base.C["orange"], 1.0)

    base.box(slide, 5.70, 3.68, 0.74, 0.50, base.C["teal_light"], base.C["teal_dark"], label="Foresight", size=6.5, bold=True)
    base.add_lock(slide, 6.32, 3.58, 0.7)
    curved_arrow(slide, 5.57, 3.94, 5.70, 3.94, base.C["orange"], 1.0)
    curved_arrow(slide, 3.72, 3.80, 5.95, 3.68, base.C["blue_dark"], 0.8)

    for i, key in enumerate(["pred1", "pred8", "pred16"]):
        base.picture_cover(slide, assets[key], 6.68 + i * 0.43, 3.73, 0.37, 0.37, base.C["teal"], 0.5)
    base.textbox(slide, "predicted tactile future", 6.58, 3.52, 1.42, 0.17, 6.0, base.C["teal_dark"], True)
    curved_arrow(slide, 6.44, 3.94, 6.66, 3.94, base.C["teal"], 1.0)

    base.token_row(slide, 8.15, 3.74, 4, base.C["orange"], 0.09, 0.02)
    base.token_row(slide, 8.15, 4.00, 4, base.C["cyan"], 0.09, 0.02, base.C["teal_dark"])
    base.box(slide, 8.68, 3.68, 0.68, 0.50, base.C["gold_light"], base.C["gold"], label="Scorer", size=6.5, bold=True)
    base.add_lock(slide, 9.24, 3.58, 0.7)
    curved_arrow(slide, 7.98, 3.92, 8.68, 3.95, base.C["teal"], 0.9)
    base.polyline(
        slide, [(5.31, 4.05), (5.31, 4.27), (8.52, 4.27), (8.52, 4.08), (8.68, 4.08)],
        base.C["orange"], 0.8,
    )
    base.box(slide, 9.60, 3.84, 0.17, 0.17, base.C["red"], base.C["red"], radius=0.04)
    base.textbox(slide, "low quality", 9.80, 3.81, 0.62, 0.20, 6.0, base.C["red"], True, PP_ALIGN.LEFT)
    curved_arrow(slide, 9.36, 3.94, 9.59, 3.94, base.C["red"], 1.0)

    # Main diffusion timeline uses fewer, larger states for legibility.
    y = 4.95
    base.draw_diffusion_strip(slide, 0.32, y, [base.C["gray_light"], base.C["gray"]], 5, 0.15, 0.025)
    base.textbox(slide, "initial noise", 0.27, 5.19, 0.92, 0.17, 6.2, base.C["muted"])
    base.connector(slide, 1.22, y + 0.08, 1.68, y + 0.08, base.C["ink"], 1.0)
    base.textbox(slide, "...", 1.47, y - 0.02, 0.25, 0.14, 8, base.C["muted"], True)

    base.box(slide, 1.82, y - 0.09, 1.25, 0.35, base.C["blue_light"], base.C["blue_dark"], radius=0.06, width=1.3)
    base.draw_diffusion_strip(slide, 1.92, y, [base.C["blue"], base.C["blue_light"]], 5, 0.14, 0.025)
    base.textbox(slide, "selected late state", 1.78, 5.20, 1.32, 0.18, 6.3, base.C["blue_dark"], True)

    base.connector(slide, 3.07, y + 0.08, 3.53, y + 0.08, base.C["green"], 1.7)
    base.textbox(slide, "guided update", 3.08, 4.72, 0.46, 0.20, 6.0, base.C["green_dark"], True)
    base.box(slide, 3.53, y - 0.09, 1.25, 0.35, base.C["green_light"], base.C["green_dark"], radius=0.06, width=1.3)
    base.draw_diffusion_strip(slide, 3.63, y, [base.C["green"], base.C["green_light"]], 5, 0.14, 0.025)
    base.textbox(slide, "updated state", 3.50, 5.20, 1.30, 0.18, 6.3, base.C["green_dark"], True)

    base.connector(slide, 4.78, y + 0.08, 5.15, y + 0.08, base.C["ink"], 1.0)
    base.draw_diffusion_strip(slide, 5.15, y, [base.C["blue_dark"]], 7, 0.15, 0.025)
    base.textbox(slide, "continued frozen denoising", 5.05, 5.20, 1.55, 0.18, 6.1, base.C["muted"], True)
    base.connector(slide, 6.42, y + 0.08, 7.00, y + 0.08, base.C["ink"], 1.0)
    base.textbox(slide, "...", 6.63, y - 0.02, 0.25, 0.14, 8, base.C["muted"], True)

    base.draw_diffusion_strip(slide, 7.05, y, [base.C["orange"]], 6, 0.15, 0.025)
    base.textbox(slide, "final guided action", 7.00, 5.20, 1.20, 0.18, 6.3, base.C["orange_dark"], True)

    execution = ASSETS / "execution_clean/chip_execution_clean_t10.0s_f0240.png"
    base.connector(slide, 8.18, y + 0.08, 8.63, y + 0.08, base.C["ink"], 1.0)
    base.picture_cover(slide, execution, 8.66, 4.72, 0.80, 0.54, base.C["green"], 0.8)
    base.textbox(slide, "execute & replan", 8.57, 5.28, 1.02, 0.17, 6.3, base.C["green_dark"], True)

    # Clean feedback routing: neither path crosses a module or prediction image.
    base.polyline(
        slide, [(9.68, 4.04), (9.68, 4.49), (2.43, 4.49), (2.43, 4.86)],
        base.C["orange"], 1.30, True,
    )
    base.textbox(slide, "quality gradient", 8.72, 4.28, 0.92, 0.17, 6.1, base.C["orange_dark"], True)
    base.polyline(
        slide, [(2.43, 4.86), (2.43, 4.43), (4.44, 4.43), (4.44, 4.18)],
        base.C["blue_dark"], 1.0,
    )


def build():
    assets = base.prepare_assets()
    prs = Presentation()
    prs.slide_width = Inches(base.SLIDE_W)
    prs.slide_height = Inches(base.SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = base.C["white"]

    base.add_section_title(slide, "OFFLINE MODEL LEARNING", 0.18, 0.08, 3.0)
    # No reference bitmap and no colored panel backgrounds: only subtle separators.
    base.connector(slide, 3.82, 0.44, 3.82, 3.10, base.C["line"], 0.65, arrow=False)
    base.connector(slide, 10.12, 0.44, 10.12, 3.10, base.C["line"], 0.65, arrow=False)
    base.connector(slide, 0.15, 3.13, 13.18, 3.13, base.C["line"], 0.85, arrow=False)

    base.draw_panel_a(slide, assets)
    base.draw_panel_b(slide, assets)
    draw_scorer(slide)
    draw_inference(slide, assets)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    shutil.copy2(OUTPUT, DESKTOP)
    print(OUTPUT)
    print(DESKTOP)


if __name__ == "__main__":
    build()
