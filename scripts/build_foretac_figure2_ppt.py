#!/usr/bin/env python3
"""Build an editable PowerPoint master for the ForeTac Figure 2 method diagram."""

from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "paper/figures/figure2_chip_assets"
GENERATED = ROOT / "paper/figures/foretac_figure2_ppt_assets"
OUTPUT = ROOT / "paper/figures/ForeTac_Figure2_editable_v2_20260812.pptx"
DESKTOP_OUTPUT = Path("/home/chenshuai/Desktop/ForeTac_Figure2_editable_v2_20260812.pptx")

SLIDE_W = 13.333
SLIDE_H = 5.657


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[:2], 16), int(value[2:4], 16), int(value[4:6], 16))


C = {
    "ink": rgb("172B3A"),
    "muted": rgb("657487"),
    "line": rgb("CCD5DF"),
    "blue": rgb("4C7FB8"),
    "blue_dark": rgb("35689F"),
    "blue_light": rgb("E8F0F8"),
    "teal": rgb("2D9FA3"),
    "teal_dark": rgb("197B82"),
    "teal_light": rgb("DDF2F1"),
    "cyan": rgb("47C6D0"),
    "orange": rgb("E78A3D"),
    "orange_dark": rgb("B96425"),
    "orange_light": rgb("FCE8D6"),
    "gold": rgb("C6A24C"),
    "gold_light": rgb("F4EDD8"),
    "green": rgb("59A968"),
    "green_dark": rgb("347B45"),
    "green_light": rgb("E2F2E4"),
    "red": rgb("D85D5D"),
    "red_light": rgb("F7E1E1"),
    "gray": rgb("8E9AA8"),
    "gray_light": rgb("EEF1F4"),
    "white": rgb("FFFFFF"),
    "panel": rgb("F7F9FB"),
}


def set_alpha(shape, transparency: int) -> None:
    if not transparency:
        return
    solid = shape._element.spPr.solidFill
    if solid is None:
        return
    color = solid[0]
    alpha = color.find(qn("a:alpha"))
    if alpha is None:
        alpha = parse_xml(f"<a:alpha {nsdecls('a')} val='100000'/>")
        color.append(alpha)
    alpha.set("val", str((100 - transparency) * 1000))


def fill(shape, color: RGBColor, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    set_alpha(shape, transparency)


def line(shape, color: RGBColor, width: float = 1.0) -> None:
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def set_dash(shape, value: str = "dash") -> None:
    ln = shape._element.spPr.ln
    old = ln.find(qn("a:prstDash"))
    if old is not None:
        ln.remove(old)
    ln.append(parse_xml(f"<a:prstDash {nsdecls('a')} val='{value}'/>"))


def arrow_end(shape, kind: str = "triangle") -> None:
    ln = shape._element.spPr.ln
    old = ln.find(qn("a:tailEnd"))
    if old is not None:
        ln.remove(old)
    ln.append(parse_xml(f"<a:tailEnd {nsdecls('a')} type='{kind}' w='sm' len='sm'/>") )


def textbox(slide, value, x, y, w, h, size=7.2, color=None, bold=False,
            align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, wrap=True):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = "TXT_" + value.replace("\n", "_")[:24]
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Pt(1)
    tf.margin_right = Pt(1)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.word_wrap = wrap
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or C["ink"]
    return shape


def box(slide, x, y, w, h, fill_color, border=None, radius=0.07, width=0.9,
        label=None, size=7.0, bold=False, text_color=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = "BOX_" + (label or "module").replace("\n", "_")[:24]
    fill(shape, fill_color)
    line(shape, border or fill_color, width)
    if label:
        tf = shape.text_frame
        tf.clear()
        tf.margin_left = Pt(2)
        tf.margin_right = Pt(2)
        tf.margin_top = Pt(1)
        tf.margin_bottom = Pt(1)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = text_color or C["ink"]
    return shape


def connector(slide, x1, y1, x2, y2, color=None, width=1.1, dashed=False, arrow=True):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.name = "FLOW_connector"
    line(shape, color or C["ink"], width)
    if dashed:
        set_dash(shape)
    if arrow:
        arrow_end(shape)
    return shape


def polyline(slide, points, color, width=1.2, dashed=False, arrow=True):
    shapes = []
    for index, (p0, p1) in enumerate(zip(points[:-1], points[1:])):
        shapes.append(connector(slide, *p0, *p1, color, width, dashed, arrow and index == len(points) - 2))
    return shapes


def add_lock(slide, x, y, scale=1.0):
    # Native-shape padlock: editable and more reliable than a font glyph.
    connector(slide, x + 0.035 * scale, y + 0.09 * scale, x + 0.035 * scale, y + 0.035 * scale,
              C["gray"], 1.0, arrow=False)
    connector(slide, x + 0.035 * scale, y + 0.035 * scale, x + 0.095 * scale, y + 0.035 * scale,
              C["gray"], 1.0, arrow=False)
    connector(slide, x + 0.095 * scale, y + 0.035 * scale, x + 0.095 * scale, y + 0.09 * scale,
              C["gray"], 1.0, arrow=False)
    box(slide, x, y + 0.08 * scale, 0.13 * scale, 0.11 * scale, C["gray"], C["gray"], radius=0.02)


def token_row(slide, x, y, count, color, size=0.11, gap=0.035, border=None):
    for i in range(count):
        box(slide, x + i * (size + gap), y, size, size, color, border or color, radius=0.025, width=0.55)


def latent_grid(slide, x, y, cell=0.09, color=None, border=None):
    color = color or C["cyan"]
    border = border or C["teal_dark"]
    for row in range(3):
        for col in range(3):
            box(slide, x + col * cell, y + row * cell, cell - 0.008, cell - 0.008,
                color, border, radius=0, width=0.35)


def picture_cover(slide, path: Path, x, y, w, h, border=C["line"], width=0.7):
    with Image.open(path) as image:
        iw, ih = image.size
    src_ratio = iw / ih
    dst_ratio = w / h
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    pic.name = "IMG_" + path.stem[:24]
    if src_ratio > dst_ratio:
        crop = (1 - dst_ratio / src_ratio) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        crop = (1 - src_ratio / dst_ratio) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    frame.name = "FRAME_" + path.stem[:24]
    frame.fill.background()
    line(frame, border, width)
    return pic


def crop_image(source: Path, target: Path, box_px) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as image:
        image.crop(box_px).save(target)
    return target


def prepare_assets() -> dict[str, Path]:
    GENERATED.mkdir(parents=True, exist_ok=True)
    ci_source = ROOT / "assets/ci_vae_reconstruction.png"
    # First decoded panel from the verified GT/reconstruction comparison.
    recon = crop_image(ci_source, GENERATED / "tacvae_real_reconstruction.png", (38, 662, 450, 1048))

    multi = ROOT / "paper/figures/foretac_multitask_qpos_preview.png"
    # Board row predicted H=1/H=8/H=16 panels; these are actual held-out model predictions.
    pred1 = crop_image(multi, GENERATED / "foresight_pred_h1.png", (592, 297, 728, 430))
    pred8 = crop_image(multi, GENERATED / "foresight_pred_h8.png", (891, 297, 1027, 430))
    pred16 = crop_image(multi, GENERATED / "foresight_pred_h16.png", (1188, 297, 1325, 430))
    return {"recon": recon, "pred1": pred1, "pred8": pred8, "pred16": pred16}


def add_section_title(slide, value, x, y, w):
    textbox(slide, value, x, y, w, 0.22, 10.8, C["ink"], True, PP_ALIGN.LEFT)


def draw_panel_a(slide, assets):
    x0, y0 = 0.18, 0.47
    textbox(slide, "(a) TacVAE Pretraining", x0, y0, 3.48, 0.24, 10.2, C["ink"], True)

    frames = [140, 163, 187]
    for i, frame in enumerate(frames):
        path = ASSETS / "marker_rdp" / f"chip_ep01_f{frame:04d}_left_rdp.png"
        picture_cover(slide, path, 0.20 + i * 0.055, 1.19 - i * 0.045, 0.60, 0.60, C["line"], 0.45)
    textbox(slide, "Marker history", 0.15, 1.83, 0.82, 0.18, 6.3, C["muted"])
    connector(slide, 0.86, 1.47, 1.02, 1.47, C["teal_dark"], 1.2)

    box(slide, 1.02, 1.20, 0.52, 0.54, C["teal_light"], C["teal_dark"], label="Tactile\nEncoder", size=6.5, bold=True)
    add_lock(slide, 1.43, 1.10, 0.72)
    connector(slide, 1.54, 1.47, 1.66, 1.47, C["teal_dark"], 1.1)

    box(slide, 1.66, 1.15, 0.39, 0.25, C["teal_light"], C["teal"], label="mu", size=6.8, bold=True)
    box(slide, 1.66, 1.52, 0.39, 0.25, C["orange_light"], C["orange"], label="log var", size=5.8, bold=True)
    connector(slide, 2.05, 1.27, 2.16, 1.44, C["teal"], 0.9)
    connector(slide, 2.05, 1.64, 2.16, 1.48, C["orange"], 0.9)
    box(slide, 2.14, 1.36, 0.18, 0.18, C["white"], C["gray"], radius=1, label="~", size=7.5, bold=True)
    connector(slide, 2.32, 1.46, 2.43, 1.46, C["teal_dark"], 1.0)
    latent_grid(slide, 2.43, 1.34, 0.082)
    textbox(slide, "Spatial latent", 2.34, 1.62, 0.48, 0.17, 5.9, C["muted"])
    connector(slide, 2.68, 1.46, 2.78, 1.46, C["teal_dark"], 1.0)
    box(slide, 2.78, 1.20, 0.48, 0.54, C["blue_light"], C["blue_dark"], label="Marker\nDecoder", size=6.2, bold=True)
    connector(slide, 3.26, 1.47, 3.36, 1.47, C["blue_dark"], 1.0)
    picture_cover(slide, assets["recon"], 3.36, 1.20, 0.38, 0.54, C["line"], 0.5)
    textbox(slide, "Recon.", 3.30, 1.78, 0.48, 0.16, 5.8, C["muted"])

    # Deployment is deliberately and exclusively branched from mu.
    polyline(slide, [(1.86, 1.15), (1.86, 0.91), (2.18, 0.91)], C["teal"], 1.0, True)
    latent_grid(slide, 2.19, 0.84, 0.052, C["teal_light"], C["teal"])
    textbox(slide, "Deployment z_t", 2.38, 0.84, 0.80, 0.18, 6.1, C["teal_dark"], True)


def draw_panel_b(slide, assets):
    textbox(slide, "(b) Action-Conditioned Tactile Foresight", 3.91, 0.47, 6.15, 0.24, 10.2, C["ink"], True)
    # Raw visual inputs.
    global_img = ASSETS / "camera_pairs/chip_ep01_f0160_global.png"
    wrist_img = ASSETS / "camera_pairs/chip_ep01_f0160_wrist.png"
    picture_cover(slide, global_img, 4.00, 0.88, 0.62, 0.43, C["line"], 0.45)
    picture_cover(slide, wrist_img, 4.00, 1.35, 0.62, 0.43, C["line"], 0.45)
    textbox(slide, "Global", 4.02, 0.72, 0.28, 0.15, 5.8, C["muted"])
    textbox(slide, "Wrist", 4.33, 0.72, 0.28, 0.15, 5.8, C["muted"])
    box(slide, 4.72, 1.02, 0.60, 0.55, C["blue_light"], C["blue_dark"], label="Visual\nEncoder", size=6.8, bold=True)
    add_lock(slide, 5.20, 0.92, 0.75)
    connector(slide, 4.62, 1.33, 4.72, 1.30, C["blue_dark"], 1.0)
    token_row(slide, 5.40, 1.23, 3, C["blue"], 0.11, 0.025)

    tactile = ASSETS / "marker_rdp/chip_ep01_f0163_left_rdp.png"
    picture_cover(slide, tactile, 4.02, 1.93, 0.62, 0.62, C["line"], 0.45)
    textbox(slide, "Touch", 4.08, 2.56, 0.50, 0.15, 5.8, C["muted"])
    box(slide, 4.72, 1.98, 0.60, 0.48, C["teal_light"], C["teal_dark"], label="TacVAE", size=7.2, bold=True)
    add_lock(slide, 5.20, 1.88, 0.75)
    connector(slide, 4.64, 2.22, 4.72, 2.22, C["teal_dark"], 1.0)
    latent_grid(slide, 5.42, 2.09, 0.075)

    token_row(slide, 4.08, 2.83, 4, C["gray_light"], 0.09, 0.025, C["gray"])
    textbox(slide, "Robot state", 4.00, 2.95, 0.58, 0.14, 5.8, C["muted"])
    box(slide, 4.72, 2.77, 0.60, 0.34, C["gray_light"], C["gray"], label="Projection", size=6.2, bold=True)
    connector(slide, 4.55, 2.88, 4.72, 2.94, C["gray"], 0.9)
    token_row(slide, 5.47, 2.87, 1, C["gray"], 0.11, 0.02)

    # Observation context rail and central core.
    polyline(slide, [(5.80, 1.29), (5.92, 1.29), (5.92, 2.88), (5.80, 2.88)], C["blue_dark"], 1.0, False, False)
    textbox(slide, "Obs.\ncontext", 5.82, 1.90, 0.45, 0.42, 6.0, C["blue_dark"], True)
    connector(slide, 5.92, 2.08, 6.25, 2.08, C["blue_dark"], 1.3)

    box(slide, 6.25, 1.43, 1.48, 1.25, C["blue_light"], C["blue_dark"], label="Foresight\nTransformer", size=9.0, bold=True)
    # Orange cross-attention band.
    box(slide, 6.31, 2.36, 1.36, 0.25, C["orange_light"], C["orange"], label="Action cross-attention", size=5.8, bold=True)

    textbox(slide, "Future queries", 6.42, 0.79, 1.10, 0.17, 6.4, C["muted"], True)
    token_row(slide, 6.48, 1.00, 3, C["gray"], 0.11, 0.04)
    textbox(slide, "...", 6.91, 0.99, 0.20, 0.12, 7.0, C["gray"], True)
    token_row(slide, 7.12, 1.00, 2, C["gray"], 0.11, 0.04)
    for x in [6.535, 6.685, 6.835, 7.175, 7.325]:
        connector(slide, x, 1.12, x, 1.43, C["gray"], 0.8)

    # Future actions are visibly separate from observation context.
    textbox(slide, "Actions", 5.84, 2.70, 0.48, 0.16, 6.1, C["orange_dark"], True)
    for i in range(3):
        # Editable small trajectory strokes.
        polyline(slide, [(5.98, 2.96 + i * 0.05), (6.11, 2.86 + i * 0.04), (6.24, 2.95 + i * 0.04)], C["orange"], 0.8, False, False)
    box(slide, 6.34, 2.77, 0.62, 0.32, C["orange_light"], C["orange"], label="Action\nEmbedding", size=5.8, bold=True)
    token_row(slide, 7.03, 2.87, 5, C["orange"], 0.10, 0.025)
    for x in [7.08, 7.205, 7.33, 7.455, 7.58]:
        connector(slide, x, 2.87, x, 2.62, C["orange"], 0.85)

    # Parallel future latent outputs and actual decoded predictions.
    connector(slide, 7.73, 2.08, 7.93, 2.08, C["teal_dark"], 1.2)
    textbox(slide, "Future tactile latents", 7.83, 0.78, 1.62, 0.18, 6.5, C["teal_dark"], True)
    latent_grid(slide, 7.97, 1.10, 0.075)
    textbox(slide, "t+1", 7.97, 1.34, 0.24, 0.14, 5.8, C["muted"])
    textbox(slide, "...", 8.33, 1.14, 0.22, 0.14, 7.0, C["muted"], True)
    latent_grid(slide, 8.59, 1.10, 0.075)
    textbox(slide, "t+H", 8.57, 1.34, 0.28, 0.14, 5.8, C["muted"])
    connector(slide, 8.22, 1.25, 8.55, 1.25, C["teal"], 0.7, False, False)

    box(slide, 8.02, 1.69, 0.66, 0.42, C["gray_light"], C["gray"], label="Marker\nDecoder", size=5.9, bold=True)
    add_lock(slide, 8.57, 1.60, 0.7)
    connector(slide, 8.34, 1.47, 8.34, 1.69, C["gray"], 0.9)
    pred_paths = [assets["pred1"], assets["pred8"], assets["pred16"]]
    for i, path in enumerate(pred_paths):
        picture_cover(slide, path, 8.86 + i * 0.39, 1.77, 0.34, 0.34, C["line"], 0.4)
    connector(slide, 8.68, 1.90, 8.84, 1.94, C["teal_dark"], 0.9)
    textbox(slide, "Decoded contact evolution", 8.72, 2.14, 1.30, 0.18, 6.2, C["muted"], True)
    textbox(slide, "parallel future", 8.75, 2.34, 1.25, 0.14, 5.5, C["gray"])


def draw_quality_icon(slide, x, y, color, symbol, label):
    box(slide, x, y, 0.34, 0.34, C["green_light"] if color == C["green"] else C["red_light"], color, radius=0.05, width=0.7)
    textbox(slide, symbol, x, y, 0.34, 0.34, 10, color, True)
    textbox(slide, label, x - 0.08, y + 0.36, 0.50, 0.18, 5.5, C["muted"])


def draw_panel_c(slide):
    textbox(slide, "(c) Contact-Quality Scorer", 10.20, 0.47, 2.95, 0.24, 10.2, C["ink"], True)
    items = [(C["green"], "+", "stable"), (C["red"], "-", "dropout"), (C["red"], "!", "excess"),
             (C["red"], "~", "slip"), (C["red"], "x", "jam")]
    for i, item in enumerate(items):
        draw_quality_icon(slide, 10.30 + i * 0.54, 0.90, *item)
    connector(slide, 10.26, 1.50, 13.05, 1.50, C["line"], 0.65, arrow=False)

    textbox(slide, "Aligned trajectories", 10.25, 1.64, 1.25, 0.18, 6.4, C["muted"], True)
    token_row(slide, 10.35, 1.93, 5, C["orange"], 0.13, 0.035)
    token_row(slide, 10.35, 2.24, 5, C["cyan"], 0.13, 0.035, C["teal_dark"])
    connector(slide, 11.15, 2.01, 11.48, 2.13, C["orange"], 1.0)
    connector(slide, 11.15, 2.31, 11.48, 2.18, C["teal"], 1.0)
    box(slide, 11.48, 1.83, 0.82, 0.68, C["gold_light"], C["gold"], label="Contact-Quality\nScorer", size=7.2, bold=True)

    textbox(slide, "stable", 12.39, 1.79, 0.48, 0.15, 5.8, C["green_dark"], True, PP_ALIGN.LEFT)
    box(slide, 12.40, 1.97, 0.55, 0.10, C["green"], C["green"], radius=0.01)
    textbox(slide, "failure risk", 12.39, 2.14, 0.62, 0.15, 5.8, C["red"], True, PP_ALIGN.LEFT)
    box(slide, 12.40, 2.32, 0.33, 0.10, C["red"], C["red"], radius=0.01)
    connector(slide, 12.30, 2.16, 12.38, 2.16, C["gold"], 1.0)


def draw_diffusion_strip(slide, x, y, colors, count=6, token=0.16, gap=0.035):
    for i in range(count):
        color = colors[min(i, len(colors) - 1)]
        box(slide, x + i * (token + gap), y, token, token, color, C["gray"] if color == C["gray_light"] else color,
            radius=0.035, width=0.55)


def draw_inference(slide, assets):
    add_section_title(slide, "INFERENCE-TIME PREDICTIVE GUIDANCE", 0.18, 3.24, 4.2)
    textbox(slide, "Frozen models; only the action sample is updated.", 0.18, 3.47, 3.1, 0.16, 6.5, C["muted"], False, PP_ALIGN.LEFT)

    # Primary timeline first so the branch reads as a local intervention.
    y = 5.00
    draw_diffusion_strip(slide, 0.28, y, [C["gray_light"], C["gray_light"], C["gray"]], 5, 0.15, 0.025)
    textbox(slide, "initial noise", 0.25, 5.21, 0.90, 0.17, 6.2, C["muted"])
    connector(slide, 1.18, y + 0.08, 1.52, y + 0.08, C["ink"], 1.0)
    textbox(slide, "...", 1.51, y - 0.02, 0.25, 0.14, 8, C["muted"], True)
    connector(slide, 1.76, y + 0.08, 2.05, y + 0.08, C["blue_dark"], 1.0)

    # Selected xk rail and update.
    box(slide, 2.05, y - 0.08, 1.18, 0.33, C["blue_light"], C["blue_dark"], radius=0.06, width=1.3)
    draw_diffusion_strip(slide, 2.14, y, [C["blue"], C["blue"], C["blue_light"]], 5, 0.14, 0.025)
    textbox(slide, "selected late state  xᵏ", 2.00, 5.29, 1.32, 0.18, 6.5, C["blue_dark"], True)
    connector(slide, 3.23, y + 0.08, 3.63, y + 0.08, C["green"], 1.6)
    textbox(slide, "guided\nupdate", 3.23, 4.69, 0.40, 0.30, 6.2, C["green_dark"], True)
    box(slide, 3.63, y - 0.08, 1.18, 0.33, C["green_light"], C["green_dark"], radius=0.06, width=1.3)
    draw_diffusion_strip(slide, 3.72, y, [C["green"], C["green"], C["green_light"]], 5, 0.14, 0.025)
    textbox(slide, "updated state  x̃ᵏ", 3.62, 5.29, 1.20, 0.18, 6.5, C["green_dark"], True)

    connector(slide, 4.81, y + 0.08, 5.20, y + 0.08, C["ink"], 1.0)
    draw_diffusion_strip(slide, 5.20, y, [C["blue_dark"]], 7, 0.15, 0.025)
    textbox(slide, "continue frozen denoising", 5.14, 5.28, 1.55, 0.18, 6.2, C["muted"], True)
    connector(slide, 6.47, y + 0.08, 6.82, y + 0.08, C["ink"], 1.0)
    textbox(slide, "...", 6.80, y - 0.02, 0.25, 0.14, 8, C["muted"], True)
    connector(slide, 7.05, y + 0.08, 7.38, y + 0.08, C["ink"], 1.0)
    draw_diffusion_strip(slide, 7.38, y, [C["orange"], C["orange"], C["green"]], 6, 0.15, 0.025)
    textbox(slide, "final guided action", 7.35, 5.28, 1.20, 0.18, 6.4, C["orange_dark"], True)

    # Real execution frames, two moments under a short time arrow.
    exec1 = ASSETS / "execution_clean/chip_execution_clean_t06.0s_f0144.png"
    exec2 = ASSETS / "execution_clean/chip_execution_clean_t10.0s_f0240.png"
    connector(slide, 8.53, y + 0.08, 9.02, y + 0.08, C["ink"], 1.0)
    picture_cover(slide, exec1, 9.03, 4.77, 0.72, 0.48, C["line"], 0.55)
    connector(slide, 9.77, 5.01, 9.97, 5.01, C["green_dark"], 0.9)
    picture_cover(slide, exec2, 9.98, 4.77, 0.72, 0.48, C["line"], 0.55)
    textbox(slide, "execute & replan", 9.13, 5.28, 1.46, 0.18, 6.4, C["green_dark"], True)

    # Compact predictive scoring branch above xk.
    token_row(slide, 2.32, 3.77, 3, C["blue"], 0.11, 0.025)
    latent_grid(slide, 2.72, 3.71, 0.065, C["teal_light"], C["teal"])
    textbox(slide, "Obs. context", 2.18, 3.52, 0.85, 0.17, 6.1, C["blue_dark"], True)

    box(slide, 3.28, 3.70, 0.72, 0.47, C["blue_light"], C["blue_dark"], label="Denoiser", size=7.0, bold=True)
    add_lock(slide, 3.88, 3.61, 0.7)
    # xk to denoiser and context to denoiser.
    polyline(slide, [(2.64, 4.92), (2.64, 4.42), (3.50, 4.42), (3.50, 4.17)], C["blue_dark"], 1.0)
    connector(slide, 3.00, 3.88, 3.28, 3.94, C["blue_dark"], 0.9)

    token_row(slide, 4.18, 3.84, 4, C["orange"], 0.12, 0.025)
    textbox(slide, "Clean action estimate", 4.05, 3.59, 0.95, 0.20, 6.1, C["orange_dark"], True)
    connector(slide, 4.00, 3.94, 4.16, 3.94, C["orange"], 1.0)

    box(slide, 4.87, 3.70, 0.72, 0.47, C["teal_light"], C["teal_dark"], label="Foresight", size=7.0, bold=True)
    add_lock(slide, 5.47, 3.61, 0.7)
    connector(slide, 4.76, 3.94, 4.87, 3.94, C["orange"], 1.0)
    polyline(slide, [(2.98, 3.88), (3.10, 3.88), (3.10, 3.48), (5.23, 3.48), (5.23, 3.70)], C["blue_dark"], 0.8)

    # Actual predicted outputs from held-out model evaluation.
    for i, key in enumerate(["pred1", "pred8", "pred16"]):
        picture_cover(slide, assets[key], 5.78 + i * 0.46, 3.71, 0.40, 0.40, C["teal"], 0.5)
    textbox(slide, "Predicted tactile future", 5.72, 3.48, 1.40, 0.17, 6.2, C["teal_dark"], True)
    connector(slide, 5.59, 3.94, 5.76, 3.94, C["teal"], 1.0)

    # Scorer receives both action and future tactile representations.
    token_row(slide, 7.24, 3.73, 4, C["orange"], 0.10, 0.02)
    token_row(slide, 7.24, 4.00, 4, C["cyan"], 0.10, 0.02, C["teal_dark"])
    connector(slide, 7.08, 3.94, 7.22, 4.05, C["teal"], 0.8)
    polyline(slide, [(4.47, 4.06), (4.47, 4.28), (7.12, 4.28), (7.12, 3.78), (7.22, 3.78)], C["orange"], 0.8)
    box(slide, 7.83, 3.70, 0.70, 0.47, C["gold_light"], C["gold"], label="Scorer", size=7.0, bold=True)
    add_lock(slide, 8.41, 3.61, 0.7)
    connector(slide, 7.72, 3.94, 7.83, 3.94, C["gold"], 1.0)
    box(slide, 8.72, 3.85, 0.18, 0.18, C["red"], C["red"], radius=0.04)
    textbox(slide, "low quality", 8.94, 3.83, 0.64, 0.20, 6.2, C["red"], True, PP_ALIGN.LEFT)
    connector(slide, 8.53, 3.94, 8.70, 3.94, C["red"], 1.0)

    # The only backward path; its endpoint is x^k, not x~^k.
    polyline(slide, [(8.82, 4.04), (8.82, 4.55), (2.64, 4.55), (2.64, 4.91)], C["orange"], 1.25, True)
    textbox(slide, "quality gradient", 7.63, 4.58, 0.90, 0.16, 6.2, C["orange_dark"], True)


def build() -> None:
    prepared = prepare_assets()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = C["white"]

    add_section_title(slide, "OFFLINE MODEL LEARNING", 0.18, 0.08, 3.0)
    # Subtle panel surfaces reproduce the reference figure's visual grouping
    # while keeping the slide background itself pure white.
    for x, w in [(0.08, 3.68), (3.88, 6.18), (10.16, 3.09)]:
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(0.42), Inches(w), Inches(2.65))
        panel.name = "PANEL_offline"
        fill(panel, C["panel"])
        panel.line.fill.background()
        slide.shapes._spTree.remove(panel._element)
        slide.shapes._spTree.insert(2, panel._element)
    infer_panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.08), Inches(3.18), Inches(13.17), Inches(2.40))
    infer_panel.name = "PANEL_inference"
    fill(infer_panel, C["panel"])
    infer_panel.line.fill.background()
    slide.shapes._spTree.remove(infer_panel._element)
    slide.shapes._spTree.insert(2, infer_panel._element)
    connector(slide, 3.82, 0.44, 3.82, 3.12, C["line"], 0.7, arrow=False)
    connector(slide, 10.12, 0.44, 10.12, 3.12, C["line"], 0.7, arrow=False)
    connector(slide, 0.15, 3.13, 13.18, 3.13, C["line"], 0.9, arrow=False)

    draw_panel_a(slide, prepared)
    draw_panel_b(slide, prepared)
    draw_panel_c(slide)
    draw_inference(slide, prepared)

    # Reference page for direct side-by-side tracing in PowerPoint.
    reference = Path("/home/chenshuai/Downloads/Generated Image August 12, 2026 - 4_10PM (1).jpg")
    if reference.exists():
        ref_slide = prs.slides.add_slide(prs.slide_layouts[6])
        ref_background = ref_slide.background.fill
        ref_background.solid()
        ref_background.fore_color.rgb = C["white"]
        picture_cover(ref_slide, reference, 0, 0, SLIDE_W, SLIDE_H, C["white"], 0)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    shutil.copy2(OUTPUT, DESKTOP_OUTPUT)
    print(OUTPUT)
    print(DESKTOP_OUTPUT)


if __name__ == "__main__":
    build()
