#!/usr/bin/env python3
"""Build a hybrid Figure 2 PPT that preserves the generated visual design."""

from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path("/home/chenshuai/Downloads/Generated Image August 12, 2026 - 4_10PM (1).jpg")
ASSETS = ROOT / "paper/figures/figure2_chip_assets"
BOARD = ROOT / "paper/figures/board_marker_displacement/panels"
PPT_ASSETS = ROOT / "paper/figures/foretac_figure2_ppt_assets"
OUTPUT = ROOT / "paper/figures/ForeTac_Figure2_hybrid_v3_20260812.pptx"
DESKTOP = Path("/home/chenshuai/Desktop/ForeTac_Figure2_hybrid_v3_20260812.pptx")

W, H = 13.333, 5.657
SX, SY = W / 6336.0, H / 2688.0


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


INK = color("172B3A")
GREEN = color("4FAE63")
RED = color("D85D5D")
WHITE = color("FFFFFF")
TEAL = color("2D9FA3")
PANEL = color("F4F6F8")


def px_box(x, y, w, h):
    return x * SX, y * SY, w * SX, h * SY


def picture_cover(
    slide, path: Path, x, y, w, h, border=WHITE, border_width=0.7,
    extra_crop_top=0.0, extra_crop_bottom=0.0,
):
    with Image.open(path) as im:
        iw, ih = im.size
    src = iw / ih
    dst = w / h
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    pic.name = "REAL_" + path.stem[:25]
    if src > dst:
        crop = (1 - dst / src) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    else:
        crop = (1 - src / dst) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    pic.crop_top = min(0.95, pic.crop_top + extra_crop_top)
    pic.crop_bottom = min(0.95, pic.crop_bottom + extra_crop_bottom)
    frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    frame.name = "EDITABLE_FRAME_" + path.stem[:20]
    frame.fill.background()
    frame.line.color.rgb = border
    frame.line.width = Pt(border_width)
    return pic


def rect(slide, x, y, w, h, fill_color, border=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border or fill_color
    return shape


def label(slide, value, x, y, w, h, size=6.5, fill=None, font_color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = "EDITABLE_TEXT_" + value[:20]
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Pt(1)
    tf.margin_right = Pt(1)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = value
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = font_color
    return shape


def badge(slide, value, x, y, fill_color):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.17), Inches(0.17))
    s.name = "EDITABLE_SAMPLE_BADGE"
    s.fill.solid()
    s.fill.fore_color.rgb = fill_color
    s.line.color.rgb = WHITE
    s.line.width = Pt(0.8)
    tf = s.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = value
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = WHITE


def add_reference(slide):
    picture_cover(slide, REFERENCE, 0, 0, W, H, WHITE, 0)


def build_hybrid_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_reference(slide)

    # (a) replace synthetic marker history and reconstruction with verified assets.
    marker_frames = [140, 163, 187]
    for i, frame in enumerate(marker_frames):
        path = ASSETS / "marker_rdp" / f"chip_ep01_f{frame:04d}_left_rdp.png"
        x, y, w, h = px_box(90 + i * 28, 525 - i * 18, 330, 330)
        picture_cover(slide, path, x, y, w, h, TEAL, 0.6)
    recon = PPT_ASSETS / "tacvae_real_reconstruction.png"
    rect(slide, *px_box(1900, 900, 410, 530), PANEL, PANEL)
    x, y, w, h = px_box(1970, 1020, 285, 285)
    picture_cover(slide, recon, x, y, w, h, TEAL, 0.6, extra_crop_top=0.09)
    label(slide, "Reconstructed\nmarker field", *px_box(1920, 1315, 370, 100), 6.1, PANEL, INK, False)

    # (b) synchronized chip global/wrist/touch observations.
    for path, box_px in [
        (ASSETS / "camera_pairs/chip_ep01_f0160_global.png", (2395, 275, 285, 210)),
        (ASSETS / "camera_pairs/chip_ep01_f0160_wrist.png", (2395, 510, 285, 210)),
        (ASSETS / "marker_rdp/chip_ep01_f0163_left_rdp.png", (2395, 790, 300, 300)),
    ]:
        picture_cover(slide, path, *px_box(*box_px), TEAL, 0.55)

    # Replace decoded future cards with genuine held-out multi-step predictions.
    preds = [PPT_ASSETS / "foresight_pred_h1.png", PPT_ASSETS / "foresight_pred_h8.png", PPT_ASSETS / "foresight_pred_h16.png"]
    for i, path in enumerate(preds):
        picture_cover(slide, path, *px_box(4135 + i * 285, 1090, 245, 245), TEAL, 0.65)

    # (c) replace generic pictograms with real positive/negative tactile examples.
    examples = [
        (BOARD / "board_marker_contact_modes_stable_contact.png", "+", GREEN),
        (ASSETS / "marker_rdp/chip_ep01_f0046_left_rdp.png", "-", RED),
        (BOARD / "board_marker_contact_modes_insufficient_pressure.png", "-", RED),
        (BOARD / "board_marker_contact_modes_excessive_pressure.png", "-", RED),
        (BOARD / "board_marker_contact_modes_oscillation.png", "-", RED),
    ]
    # A self-contained example strip replaces the generated icons and labels.
    rect(slide, *px_box(4860, 300, 1425, 465), PANEL, PANEL)
    label(slide, "Contact-quality examples", *px_box(5100, 320, 970, 65), 7.2, PANEL, INK, True)
    for i, (path, sign, sign_color) in enumerate(examples):
        x, y, w, h = px_box(4950 + i * 265, 400, 205, 185)
        picture_cover(slide, path, x, y, w, h, sign_color, 1.3)
        badge(slide, sign, x + w - 0.13, y - 0.03, sign_color)
    labels = ["stable", "dropout", "low force", "excess", "oscillation"]
    for i, value in enumerate(labels):
        x, y, w, h = px_box(4915 + i * 265, 602, 270, 65)
        label(slide, value, x, y, w, h, 5.6, PANEL, INK, i == 0)

    # Use one strong real execution frame in the original terminal-result slot.
    rect(slide, *px_box(5840, 2040, 440, 610), PANEL, PANEL)
    picture_cover(
        slide,
        ASSETS / "execution_clean/chip_execution_clean_t10.0s_f0240.png",
        *px_box(5910, 2120, 310, 310), GREEN, 0.9,
    )
    label(slide, "Execute & replan", *px_box(5890, 2460, 380, 80), 6.4, PANEL, GREEN, True)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    build_hybrid_slide(prs)
    ref = prs.slides.add_slide(prs.slide_layouts[6])
    add_reference(ref)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    shutil.copy2(OUTPUT, DESKTOP)
    print(OUTPUT)
    print(DESKTOP)


if __name__ == "__main__":
    main()
