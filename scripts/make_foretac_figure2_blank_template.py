#!/usr/bin/env python3
"""Create a text-free, image-free editable Figure 2 template from redrawn v4."""

from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper/figures/ForeTac_Figure2_redrawn_v4_20260812.pptx"
OUTPUT = ROOT / "paper/figures/ForeTac_Figure2_blank_pastel_editable_20260812.pptx"


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def clear_text(shape) -> None:
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.text = ""
    shape.text_frame.clear()


def main() -> None:
    prs = Presentation(SOURCE)
    for slide in prs.slides:
        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        for picture in pictures:
            placeholder = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                picture.left,
                picture.top,
                picture.width,
                picture.height,
            )
            placeholder.fill.solid()
            placeholder.fill.fore_color.rgb = RGBColor(255, 255, 255)
            placeholder.line.color.rgb = RGBColor(205, 213, 219)
            placeholder.line.width = 6350
            picture._element.addprevious(deepcopy(placeholder._element))
            remove_shape(placeholder)
            remove_shape(picture)

        for shape in slide.shapes:
            clear_text(shape)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
